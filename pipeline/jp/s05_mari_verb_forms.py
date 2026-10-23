"""Source 5: Mari's "Verb Forms Summary" tab (same workbook as source 4).
Run: python s05_mari_verb_forms.py <path to the exported .xlsx>

Shape: a visual reference laid out in cells, not a table. A verb is split across cells
(D=あ, E=う -> あう) and its result is too (O=あ, P=って -> あって). Rule text sits in
between (～　って), and the long/short form tables carry a romaji line under each kana
line.

What this source is FOR: it is the gold set for conjugate.py. Every (verb, form) pair
here is a known-correct answer from the tutor, so the engine that will generate drill
items is tested against it before it generates anything. That turns the OPM's
Evaluating step into a deterministic check for this whole family of items.

It also checks the sheet: the printed romaji is compared with a romanization of the
kana, the same trick as the vocabulary extractor, and a mismatch is reported.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

from common import bundle, ingest, is_kana, nfc, write
from conjugate import conjugate

SOURCE_ID = "s05_mari_verb_forms"

SECTIONS = {
    "Masu-form": "masu", "Te-form": "te", "Nai-form": "nai", "Ta-form": "ta",
    "Potential-form": "potential", "Volitional-form": "volitional",
    "Conditional-form": "conditional", "Imperative-form": "imperative",
    "Negative Imperative-form": "prohibitive", "Passive-form": "passive",
    "Causative-form": "causative", "Causative Passive-form": "causative_passive",
    "Long form & Short form (Present tense)": "long_short_present",
    "Long form & Short form (Past tense)": "long_short_past",
}
LEFT, RIGHT = "CDE", "MNOP"  # dictionary form cells, result cells
LONG_SHORT = {"B": "long_affirmative", "E": "long_negative",
              "I": "short_affirmative", "L": "short_negative"}

# Minimal Hepburn, enough for the forms on this sheet. Long vowels are written out
# (ou, ii), which is how the sheet spells them.
KANA = {
    "きゃ": "kya", "きゅ": "kyu", "きょ": "kyo", "しゃ": "sha", "しゅ": "shu", "しょ": "sho",
    "ちゃ": "cha", "ちゅ": "chu", "ちょ": "cho", "にゃ": "nya", "にゅ": "nyu", "にょ": "nyo",
    "りゃ": "rya", "りゅ": "ryu", "りょ": "ryo", "じゃ": "ja", "じゅ": "ju", "じょ": "jo",
}
BASE = dict(zip(
    "あいうえおかきくけこさしすせそたちつてとなにぬねのはひふへほまみむめもやゆよらりるれろわをん"
    "がぎぐげござじずぜぞだぢづでどばびぶべぼぱぴぷぺぽ",
    "a i u e o ka ki ku ke ko sa shi su se so ta chi tsu te to na ni nu ne no ha hi fu he ho "
    "ma mi mu me mo ya yu yo ra ri ru re ro wa o n ga gi gu ge go za ji zu ze zo da ji zu de do "
    "ba bi bu be bo pa pi pu pe po".split()))


def romanize(kana: str) -> str:
    out, i = "", 0
    while i < len(kana):
        if kana[i:i + 2] in KANA:
            out, i = out + KANA[kana[i:i + 2]], i + 2
        elif kana[i] == "っ" and i + 1 < len(kana):
            nxt = KANA.get(kana[i + 1:i + 3]) or BASE.get(kana[i + 1], "")
            out, i = out + ("t" if nxt.startswith("ch") else nxt[:1]), i + 1
        else:
            out, i = out + BASE.get(kana[i], kana[i]), i + 1
    return out


def _clean(v) -> str:
    s = nfc(str(v)) if v is not None else ""
    s = re.sub(r"\([A-Za-z\s]+\)?", "", s)  # やすむ(Yasumu) -> やすむ
    return s.replace("※", "").replace(" ", "").replace("　", "")


def run(xlsx: Path) -> dict:
    raw = ingest(xlsx, SOURCE_ID)
    ws = openpyxl.load_workbook(xlsx)["Verb Forms Summary"]

    section, grp, items, notes, failures = None, None, [], [], []
    pending_ls = None  # a long/short kana row waiting for its romaji row
    for r in range(21, ws.max_row + 1):
        a = _clean(ws[f"A{r}"].value)
        title = nfc(str(ws[f"A{r}"].value or ""))
        if title in SECTIONS:
            section, grp = SECTIONS[title], None
            continue
        m = re.match(r"Group(\d)", a)
        if m:
            grp = int(m.group(1))
            continue
        if title in ("I-adjective", "Na-adjective", "Noun"):
            grp = title
            continue
        if section is None:
            continue

        if section.startswith("long_short"):
            row = {k: _clean(ws[f"{c}{r}"].value) for c, k in LONG_SHORT.items()}
            if not row["long_affirmative"] or row["long_affirmative"] in ("Longform", "Affirmative"):
                continue
            if is_kana(row["long_affirmative"]):
                pending_ls = {"id": f"{SOURCE_ID}.r{r}", "section": section, "group": grp,
                              **row}
                items.append(pending_ls)
            elif pending_ls:
                for k in LONG_SHORT.values():
                    printed, own = row[k].lower(), romanize(pending_ls[k])
                    if printed and printed != own:
                        notes.append(f"romaji mismatch, row {r}: {pending_ls[k]} is printed "
                                     f"{row[k]!r}, should be {own!r} (romaji is not kept)")
                pending_ls = None
            continue

        left = "".join(_clean(ws[f"{c}{r}"].value) for c in LEFT)
        right = "".join(_clean(ws[f"{c}{r}"].value) for c in RIGHT)
        # Some rows put the stem one column to the left (L instead of M).
        if right and not right.startswith(left[:1]) and ws[f"L{r}"].value:
            right = _clean(ws[f"L{r}"].value) + right
        if not left or "Dictionary" in left or not right or not is_kana(left):
            continue
        items.append({"id": f"{SOURCE_ID}.r{r}", "section": section, "group": grp,
                      "dictionary": left, "result": right})
    write(SOURCE_ID, "mined", {"items": items})

    # ── Validating: the engine against Mari, cell by cell.
    agree = 0
    for it in items:
        if it["section"].startswith("long_short"):
            if not isinstance(it["group"], int):
                continue  # adjectives and nouns: gold data, no engine yet
            dic = it["short_affirmative"]
            past = it["section"].endswith("past")
            c = conjugate(dic if not past else _dic_from_past(it), it["group"])
            want = {"long_affirmative": c["masu"][:-2] + ("ました" if past else "ます"),
                    "long_negative": c["masu"][:-2] + ("ませんでした" if past else "ません"),
                    "short_affirmative": c["ta"] if past else (dic if not past else ""),
                    "short_negative": c["nai"][:-1] + ("かった" if past else "い")}
            for k, v in want.items():
                if it[k] != v:
                    failures.append({"item_id": it["id"], "check": "engine_matches_mari",
                                     "detail": f"{k}: sheet {it[k]}, engine {v}"})
                else:
                    agree += 1
            continue
        c = conjugate(it["dictionary"], it["group"])
        if c[it["section"]] != it["result"]:
            failures.append({"item_id": it["id"], "check": "engine_matches_mari",
                             "detail": f"{it['dictionary']} {it['section']}: "
                                       f"sheet {it['result']}, engine {c[it['section']]}"})
        else:
            agree += 1
    notes.insert(0, f"engine agrees with Mari on {agree} conjugations")
    return bundle(SOURCE_ID, raw, items, failures, notes)


PAST_TO_DIC = {"いった": "いく", "かった": "かう", "あった": "ある", "おきた": "おきる",
               "した": "する", "きた": "くる"}


def _dic_from_past(it) -> str:
    """The past table prints no dictionary form. Its verbs are the present table's."""
    return PAST_TO_DIC[it["short_affirmative"]]


if __name__ == "__main__":
    b = run(Path(sys.argv[1]))
    print(b["state"], b["item_count"], "items")
    for f in b["validation"]["failures"]:
        print("FAIL", f)
    for n in b["review_notes"]:
        print("NOTE", n)
